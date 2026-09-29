"""Build the payment gateway specification pair (v0 = flawed design, v1 = fixed design).

Both workbooks follow the 10-sheet template of the course (Hướng dẫn tạo Spec file dạng Excel).
The two versions share requirements, properties and test cases; they differ only in the
state model (MoHinhTrangThai) and in the bugs declared in HamSinhMa, so running the pipeline
on both gives a before/after comparison at the design level.

    uv run python scripts/build_payment_specs.py
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

OUT_DIR = Path(__file__).resolve().parent.parent / "dataset" / "specs_excel"

DESCRIPTION = (
    "Cổng thanh toán cho đơn hàng thương mại điện tử: vòng đời thanh toán (tạo đơn, chờ cổng, timeout, "
    "đối soát, thu tiền, giao hàng, hoàn tiền); khóa sổ cái dùng chung giữa worker xử lý webhook và worker "
    "đối soát; các hàm phân tích chữ ký webhook/mã đơn, ghi audit log, đóng phiên, giải phóng token thẻ; "
    "bộ đếm tổng tiền và giao dịch cần kiểm tra tràn số."
)

REQUIREMENTS = [
    ("REQ-1", "Chỉ được giao hàng (shipped) khi đơn đã được xác nhận thu tiền (paid)", "safety", "HIGH"),
    ("REQ-2", ("Khi cổng thanh toán timeout, đơn cuối cùng phải được đối soát xong (paid hoặc failed), "
               "không chờ vô hạn"), "liveness", "HIGH"),
    ("REQ-3", "Worker webhook và worker đối soát không được đồng thời giữ khóa sổ cái (mutual exclusion)",
     "safety", "HIGH"),
    ("REQ-4", "Hệ thống không rơi vào deadlock khi tranh chấp khóa sổ cái hoặc khi hoàn tiền", "liveness", "HIGH"),
    ("REQ-5", "Hàm parse_webhook_signature không được gây tràn bộ đệm khi nhận chữ ký từ webhook",
     "memory", "HIGH"),
    ("REQ-6", "Hàm write_audit_log phải giải phóng bộ nhớ đã cấp phát cho mỗi bản ghi audit", "memory", "MEDIUM"),
    ("REQ-7", "Hàm close_payment_session không được giải phóng cùng một con trỏ nhiều hơn một lần",
     "memory", "HIGH"),
    ("REQ-8", "Hàm release_card_token_buffer không được truy cập bộ đệm token sau khi đã giải phóng",
     "memory", "HIGH"),
    ("REQ-9", "Hàm calculate_total_payment không được để tổng tiền hàng + phí vận chuyển tràn int32",
     "safety", "HIGH"),
    ("REQ-10", "Hàm refund_ratio không được lỗi khi tổng tiền đơn hàng bằng 0", "functional", "MEDIUM"),
    ("REQ-11", "Không tồn tại mã bí mật hardcode cho phép vượt qua xác thực webhook hoặc PIN hỗ trợ",
     "safety", "HIGH"),
    ("REQ-12", "Một đơn hàng không được ghi nhận thu tiền hai lần khi cổng gửi lại webhook (idempotency)",
     "safety", "HIGH"),
    ("REQ-13", "Mọi yêu cầu hoàn tiền cuối cùng phải hoàn tất (refunded)", "liveness", "MEDIUM"),
    ("REQ-14", "Hàm parse_order_id chỉ chấp nhận mã đơn không vượt quá độ dài quy định", "memory", "MEDIUM"),
]

NFRS = [
    ("NFR-1", "Xử lý tối thiểu 500 giao dịch thanh toán/giây trên phần cứng tham chiếu", "MEDIUM"),
    ("NFR-2", "Độ trễ p95 xử lý một webhook dưới 300 ms", "MEDIUM"),
    ("NFR-3", "Audit log ghi đồng bộ, không mất bản ghi kể cả khi worker crash giữa chừng", "HIGH"),
    ("NFR-4", "Cổng gửi lại webhook tối đa 5 lần trong 24 giờ; hệ thống phải xử lý trùng lặp an toàn", "HIGH"),
    ("NFR-5", "Tính sẵn sàng 99,9% theo tháng", "MEDIUM"),
    ("NFR-6", "Không lưu số thẻ (PAN) ở dạng rõ; chỉ lưu token thẻ (PCI DSS)", "HIGH"),
]

CONSTRAINTS = [
    ("CON-1", "Chữ ký webhook là chuỗi hex 64 ký tự (HMAC-SHA256); bộ đệm 65 byte gồm ký tự kết thúc", "HIGH"),
    ("CON-2", "Mã đơn hàng tối đa 20 ký tự ASCII; bộ đệm 21 byte", "HIGH"),
    ("CON-3", "Token thẻ tối đa 32 ký tự; bộ đệm 33 byte", "HIGH"),
    ("CON-4", "Số tiền tính bằng VND, kiểu int có dấu 32-bit", "HIGH"),
    ("CON-5", "Khóa sổ cái là mutex nhẹ (futex), không dùng khóa cấp hệ điều hành nặng", "MEDIUM"),
    ("CON-6", "Mã PIN hỗ trợ viên là cặp số nguyên 16-bit không dấu", "MEDIUM"),
]

# (state, initial, labels, successors) -------------------------------------------------
LOCK_V0 = [
    ("ledger_free", True, "safe", "webhook_holds, recon_holds"),
    ("webhook_holds", False, "webhook_locked", "ledger_free, lock_conflict"),
    ("recon_holds", False, "recon_locked", "ledger_free"),
    ("lock_conflict", False, "webhook_locked, recon_locked, violation", "ledger_free"),
]
LOCK_V1 = [
    ("ledger_free", True, "safe", "webhook_holds, recon_holds"),
    ("webhook_holds", False, "webhook_locked", "ledger_free"),
    ("recon_holds", False, "recon_locked", "ledger_free"),
]
PAYMENT_V0 = [
    ("CREATED", True, "created", "PENDING"),
    ("PENDING", False, "processing", "SUCCEEDED, FAILED, UNKNOWN"),
    ("UNKNOWN", False, "timeout", "UNKNOWN, RECONCILING, SHIPPED_UNPAID"),
    ("RECONCILING", False, "reconciling", "SUCCEEDED, FAILED"),
    ("SUCCEEDED", False, "paid", "SHIPPED, DUPLICATE_CAPTURE, REFUNDING"),
    ("DUPLICATE_CAPTURE", False, "paid, double_capture", "SHIPPED"),
    ("FAILED", False, "failed", "PENDING, CREATED"),
    ("SHIPPED", False, "shipped, paid", "SHIPPED"),
    ("SHIPPED_UNPAID", False, "shipped, unpaid_shipping", "SHIPPED_UNPAID"),
    ("REFUNDING", False, "refunding", ""),
]
PAYMENT_V1 = [
    ("CREATED", True, "created", "PENDING"),
    ("PENDING", False, "processing", "SUCCEEDED, FAILED, UNKNOWN"),
    ("UNKNOWN", False, "timeout", "RECONCILING"),
    ("RECONCILING", False, "reconciling", "SUCCEEDED, FAILED"),
    ("SUCCEEDED", False, "paid", "SHIPPED, REFUNDING"),
    ("FAILED", False, "failed", "PENDING, CREATED"),
    ("SHIPPED", False, "shipped, paid", "SHIPPED"),
    ("REFUNDING", False, "refunding", "REFUNDED"),
    ("REFUNDED", False, "refunded", "REFUNDED"),
]

PROPERTIES = [
    "AG(!violation)",
    "safe -> EF(webhook_locked)",
    "AG(shipped -> paid)",
    "AG(timeout -> AF(paid | failed))",
    "AG(!double_capture)",
    "AG(!(processing & paid))",
    "created -> EF(shipped)",
    "AG(refunding -> AF(refunded))",
]

MUTEX = [
    ("webhook_locked", "recon_locked"),
    ("violation", "safe"),
    ("processing", "paid"),
    ("paid", "failed"),
    ("timeout", "paid"),
]

# (name, language, kind, params, buffer, bug_v0, magic1, magic2) --------------------------
FUNCTIONS = [
    ("parse_webhook_signature", "c", "buffer_copy", "raw_signature", 65, "buffer-overflow", None, None),
    ("parse_order_id", "c", "buffer_copy", "raw_order_id", 21, None, None, None),
    ("write_audit_log", "c", "alloc_release", "", 64, "memory-leak", None, None),
    ("close_payment_session", "c", "alloc_release", "", 16, "double-free", None, None),
    ("release_card_token_buffer", "c", "alloc_release", "", 33, "use-after-free", None, None),
    ("calculate_total_payment", "c", "arithmetic", "amount, shipping_fee", None, "integer-overflow", None, None),
    ("refund_ratio", "python", "divide", "refund_amount, total_amount", None, "divide-by-zero", None, None),
    ("webhook_bypass_token", "python", "magic_value", "header_token, secret_token", None, "hardcoded-secret",
     7788, 9999),
    ("support_override_pin", "python", "magic_value", "pin_a, pin_b", None, "hardcoded-secret", 4821, 7306),
]

ARITHMETIC = [
    ("order_total_counter", "amount, shipping_fee",
     "amount >= 0 and shipping_fee >= 0 and amount + shipping_fee > 2000000000", "amount + shipping_fee"),
    ("daily_transaction_counter", "tx_today, tx_backlog",
     "tx_today >= 0 and tx_backlog >= 0 and tx_today + tx_backlog > 32767", "tx_today + tx_backlog"),
    ("active_session_counter", "active_sessions, pending_sessions",
     "active_sessions >= 0 and pending_sessions >= 0 and active_sessions + pending_sessions > 65535",
     "active_sessions + pending_sessions"),
    ("refund_total_vnd", "refunded_vnd, pending_refund_vnd",
     "refunded_vnd >= 0 and pending_refund_vnd >= 0 and refunded_vnd + pending_refund_vnd > 2000000000",
     "refunded_vnd + pending_refund_vnd"),
    ("webhook_retry_counter", "retries_today, retries_backlog",
     "retries_today >= 0 and retries_backlog >= 0 and retries_today + retries_backlog > 32767",
     "retries_today + retries_backlog"),
    ("settlement_batch_bytes", "batch_bytes, header_bytes",
     "batch_bytes >= 0 and header_bytes >= 0 and batch_bytes + header_bytes > 2000000000",
     "batch_bytes + header_bytes"),
]

TEST_CASES = [
    ("TC-1", "REQ-3, REQ-4", "Worker webhook và worker đối soát tranh chấp khóa sổ cái",
     "Mô hình có cụm khóa ledger_free/webhook_holds/recon_holds",
     "Dựng Kripke từ MoHinhTrangThai\nKiểm tra AG(!violation)\nKiểm tra deadlock",
     "AG(!violation) ĐÚNG và không có trạng thái deadlock", "HIGH"),
    ("TC-2", "REQ-1", "Không giao hàng khi chưa thu tiền",
     "Đơn ở trạng thái UNKNOWN (cổng timeout)",
     "Kiểm tra AG(shipped -> paid)\nNếu sai, đọc phản ví dụ",
     "AG(shipped -> paid) ĐÚNG", "HIGH"),
    ("TC-3", "REQ-2", "Timeout luôn được đối soát",
     "Cổng không phản hồi", "Kiểm tra AG(timeout -> AF(paid | failed))",
     "AG(timeout -> AF(paid | failed)) ĐÚNG", "HIGH"),
    ("TC-4", "REQ-12, NFR-4", "Cổng gửi lại webhook thu tiền",
     "Đơn đã SUCCEEDED", "Kiểm tra AG(!double_capture)",
     "AG(!double_capture) ĐÚNG", "HIGH"),
    ("TC-5", "REQ-13", "Hoàn tiền luôn hoàn tất",
     "Đơn đã SUCCEEDED, khách yêu cầu hoàn tiền", "Kiểm tra AG(refunding -> AF(refunded))",
     "AG(refunding -> AF(refunded)) ĐÚNG và không có trạng thái deadlock", "MEDIUM"),
    ("TC-6", "REQ-5, CON-1", "Chữ ký webhook dài bất thường",
     "Hàm parse_webhook_signature, bộ đệm 65 byte",
     "Sinh mã C\nQuét tĩnh\nBiên dịch với ASan, gọi với chuỗi 520 ký tự",
     "Không phát hiện CWE-120 tại parse_webhook_signature", "HIGH"),
    ("TC-7", "REQ-6, REQ-7, REQ-8", "Vòng đời bộ nhớ audit log / phiên / token thẻ",
     "Hàm write_audit_log, close_payment_session, release_card_token_buffer",
     "Sinh mã C\nQuét tĩnh\nChạy với ASan",
     "Không phát hiện CWE-401, CWE-415, CWE-416", "HIGH"),
    ("TC-8", "REQ-9, CON-4", "Tổng tiền vượt int32",
     "Hàm calculate_total_payment", "Sinh mã C\nBiên dịch với UBSan, gọi với (INT_MAX, 1)",
     "Không phát hiện CWE-190 tại calculate_total_payment", "HIGH"),
    ("TC-9", "REQ-9", "Chứng minh tiền điều kiện chống tràn bằng Z3",
     "Kiểm tra số học order_total_counter",
     "Giải điều kiện tràn int32 bằng Z3\nThêm tiền điều kiện của bản vá, giải lại",
     "Z3 SAT cho order_total_counter khi chưa kiểm tra biên và UNSAT sau khi thêm tiền điều kiện", "MEDIUM"),
    ("TC-10", "REQ-10", "Tỷ lệ hoàn tiền với tổng tiền bằng 0",
     "Hàm refund_ratio", "Chạy black-box, coverage-guided và concolic fuzzing",
     "Không crash (ZeroDivisionError) tại refund_ratio", "MEDIUM"),
    ("TC-11", "REQ-11, CON-6", "Tìm mã bí mật bằng concolic fuzzing",
     "Hàm webhook_bypass_token, support_override_pin", "Chạy 3 chiến lược fuzzing với cùng ngân sách",
     "Không phát hiện CWE-798", "HIGH"),
    ("TC-12", "REQ-14, CON-2", "Mã đơn hợp lệ không bị báo lỗi nhầm",
     "Hàm parse_order_id sinh không có lỗi", "Sinh mã C\nQuét tĩnh\nChạy với ASan",
     "Không phát hiện CWE-120 tại parse_order_id", "MEDIUM"),
]


def _sheet(wb, title: str, header: list[str], rows: list[tuple]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for row in rows:
        ws.append(["" if v is None else v for v in row])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), 70)
        for c in col:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def build(version: str) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    name = f"payment_gateway_{version}"
    note = " Phiên bản v0: thiết kế ban đầu còn lỗi." if version == "v0" else " Phiên bản v1: thiết kế đã sửa."
    _sheet(wb, "ThongTin", ["Truong", "GiaTri"], [("name", name), ("description", DESCRIPTION + note)])
    _sheet(wb, "YeuCauChucNang", ["ID", "MoTa", "DanhMuc", "MucUuTien"], REQUIREMENTS)
    _sheet(wb, "YeuCauPhiChucNang", ["ID", "MoTa", "MucUuTien"], NFRS)
    _sheet(wb, "RangBuoc", ["ID", "MoTa", "MucUuTien"], CONSTRAINTS)
    states = (LOCK_V0 + PAYMENT_V0) if version == "v0" else (LOCK_V1 + PAYMENT_V1)
    _sheet(wb, "MoHinhTrangThai", ["TrangThai", "KhoiTao", "Nhan", "ChuyenDenCacTrangThai"],
           [(s, "x" if init else "", labels, nxt) for s, init, labels, nxt in states])
    _sheet(wb, "ThuocTinhAnToan", ["CongThucCTL"], [(p,) for p in PROPERTIES])
    _sheet(wb, "LoaiTruLanNhau", ["MenhDe1", "MenhDe2"], MUTEX)
    # v1 declares every function as patched: C functions with an empty Loi, Python functions with
    # Loi=safe (an empty Loi keeps the course tool's "vulnerable by construction" meaning for Python).
    funcs = []
    for fname, lang, kind, params, buf, bug, m1, m2 in FUNCTIONS:
        declared = bug if version == "v0" else ("safe" if lang == "python" else None)
        funcs.append((fname, lang, kind, params, buf, declared, m1, m2))
    _sheet(wb, "HamSinhMa", ["Ten", "NgonNgu", "Loai", "ThamSo", "KichThuocBuffer", "Loi",
                             "GiaTriBiMat1", "GiaTriBiMat2"], funcs)
    _sheet(wb, "KiemTraSoHoc", ["Ten", "BienSoNguyen", "DieuKien", "BieuThucTran"], ARITHMETIC)
    _sheet(wb, "TestCase", ["ID", "YeuCauLienQuan", "TieuDe", "DieuKienTienQuyet", "CacBuoc",
                            "KetQuaMongDoi", "MucUuTien"], TEST_CASES)
    path = OUT_DIR / f"{name}_spec.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    for v in ("v0", "v1"):
        print(f"Đã tạo {build(v)}")
