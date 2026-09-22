"""Specification parser and data models for JSON and Excel files."""
from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


@dataclass
class Requirement:
    id: str
    text: str
    category: str = "functional"  # safety, liveness, memory, functional
    priority: str = "HIGH"  # HIGH, MEDIUM, LOW


@dataclass
class StateMachine:
    states: list[str] = field(default_factory=list)
    initial: list[str] = field(default_factory=list)
    labeling: dict[str, list[str]] = field(default_factory=dict)
    transitions: list[list[str]] = field(default_factory=list)  # [[from, to], ...]


@dataclass
class FunctionSpec:
    name: str
    language: str = "c"  # "c" or "python"
    kind: str = "buffer_copy"  # buffer_copy, alloc_release, arithmetic, divide, magic_value
    params: list[str] = field(default_factory=list)
    buffer_size: int | None = None
    bug: str | None = None  # buffer-overflow, memory-leak, double-free, use-after-free, None
    magic_values: list[int] = field(default_factory=list)


@dataclass
class ArithmeticCheck:
    name: str
    int_vars: list[str] = field(default_factory=list)
    condition: str | None = None
    overflow_expr: str | None = None


@dataclass
class TestCase:
    id: str
    requirement_ids: list[str] = field(default_factory=list)
    title: str = ""
    preconditions: str = ""
    steps: list[str] = field(default_factory=list)
    expected_result: str = ""
    priority: str = "HIGH"


@dataclass
class Specification:
    name: str
    description: str = ""
    requirements: list[Requirement] = field(default_factory=list)
    state_machine: StateMachine | None = None
    additional_flows: dict[str, StateMachine] = field(default_factory=dict)
    safety_properties: list[str] = field(default_factory=list)
    mutual_exclusion: dict[str, str] | None = None
    functions: list[FunctionSpec] = field(default_factory=list)
    arithmetic_checks: list[ArithmeticCheck] = field(default_factory=list)
    test_cases: list[TestCase] = field(default_factory=list)
    source_file: str = ""

    def summary(self) -> str:
        num_states = len(self.state_machine.states) if self.state_machine else 0
        num_transitions = len(self.state_machine.transitions) if self.state_machine else 0
        return (
            f"Đặc tả '{self.name}' — {len(self.requirements)} yêu cầu, "
            f"{num_states} trạng thái ({num_transitions} chuyển tiếp), "
            f"{len(self.functions)} hàm, {len(self.arithmetic_checks)} kiểm tra số học, "
            f"{len(self.test_cases)} test case."
        )


def _parse_sm_dict(data: dict[str, Any] | None) -> StateMachine | None:
    if not data:
        return None
    return StateMachine(
        states=list(data.get("states", [])),
        initial=list(data.get("initial", [])),
        labeling={str(k): list(v) for k, v in data.get("labeling", {}).items()},
        transitions=[list(t) for t in data.get("transitions", [])],
    )


def parse_specification_json(filepath: str | Path) -> Specification:
    path = Path(filepath)
    with open(path, "r", encoding="utf-8") as f:
        raw = json.load(f)

    reqs = [
        Requirement(
            id=str(r.get("id", f"REQ-{i+1}")),
            text=str(r.get("text", "")),
            category=str(r.get("category", "functional")),
            priority=str(r.get("priority", "HIGH")),
        )
        for i, r in enumerate(raw.get("requirements", []))
    ]

    sm = _parse_sm_dict(raw.get("state_machine"))

    flows = {}
    for flow_name, flow_dict in raw.get("additional_flows", {}).items():
        flow_sm = _parse_sm_dict(flow_dict)
        if flow_sm:
            flows[flow_name] = flow_sm

    funcs = []
    for fn in raw.get("functions", []):
        funcs.append(
            FunctionSpec(
                name=str(fn.get("name", "func")),
                language=str(fn.get("language", "c")).lower(),
                kind=str(fn.get("kind", "buffer_copy")),
                params=list(fn.get("params", [])),
                buffer_size=fn.get("buffer_size"),
                bug=fn.get("bug"),
                magic_values=list(fn.get("magic_values", [])),
            )
        )

    checks = []
    for ac in raw.get("arithmetic_checks", []):
        checks.append(
            ArithmeticCheck(
                name=str(ac.get("name", "check")),
                int_vars=list(ac.get("int_vars", [])),
                condition=ac.get("condition"),
                overflow_expr=ac.get("overflow_expr"),
            )
        )

    tests = []
    for tc in raw.get("test_cases", []):
        steps_val = tc.get("steps", [])
        if isinstance(steps_val, str):
            steps = [s.strip() for s in steps_val.split("\n") if s.strip()]
        else:
            steps = [str(s) for s in steps_val]
        tests.append(
            TestCase(
                id=str(tc.get("id", f"TC-{len(tests)+1}")),
                requirement_ids=list(tc.get("requirement_ids", [])),
                title=str(tc.get("title", "")),
                preconditions=str(tc.get("preconditions", "")),
                steps=steps,
                expected_result=str(tc.get("expected_result", "")),
                priority=str(tc.get("priority", "HIGH")),
            )
        )

    return Specification(
        name=str(raw.get("name", path.stem)),
        description=str(raw.get("description", "")),
        requirements=reqs,
        state_machine=sm,
        additional_flows=flows,
        safety_properties=list(raw.get("safety_properties", [])),
        mutual_exclusion=raw.get("mutual_exclusion"),
        functions=funcs,
        arithmetic_checks=checks,
        test_cases=tests,
        source_file=str(path),
    )


def parse_specification_excel(filepath: str | Path) -> Specification:
    import openpyxl

    path = Path(filepath)
    wb = openpyxl.load_workbook(path, data_only=True)

    name = path.stem
    description = ""
    if "ThongTin" in wb.sheetnames:
        for row in wb["ThongTin"].iter_rows(values_only=True):
            if row and len(row) >= 2 and row[0]:
                k = str(row[0]).strip()
                v = str(row[1]).strip() if row[1] is not None else ""
                if k == "name":
                    name = v
                elif k == "description":
                    description = v

    reqs: list[Requirement] = []
    if "YeuCauChucNang" in wb.sheetnames:
        rows = list(wb["YeuCauChucNang"].iter_rows(values_only=True))
        if len(rows) > 1:
            for r in rows[1:]:
                if not r or not r[0]:
                    continue
                reqs.append(
                    Requirement(
                        id=str(r[0]).strip(),
                        text=str(r[1]).strip() if len(r) > 1 and r[1] else "",
                        category=str(r[2]).strip() if len(r) > 2 and r[2] else "functional",
                        priority=str(r[3]).strip() if len(r) > 3 and r[3] else "HIGH",
                    )
                )

    sm: StateMachine | None = None
    if "MoHinhTrangThai" in wb.sheetnames:
        rows = list(wb["MoHinhTrangThai"].iter_rows(values_only=True))
        if len(rows) > 1:
            states = []
            initial = []
            labeling = {}
            transitions = []
            for r in rows[1:]:
                if not r or not r[0]:
                    continue
                st = str(r[0]).strip()
                states.append(st)
                is_init = str(r[1]).strip().lower() if len(r) > 1 and r[1] is not None else ""
                if is_init in ("x", "1", "true", "yes"):
                    initial.append(st)
                labels_str = str(r[2]).strip() if len(r) > 2 and r[2] else ""
                labeling[st] = [l.strip() for l in labels_str.split(",") if l.strip()]
                target_str = str(r[3]).strip() if len(r) > 3 and r[3] else ""
                for tgt in target_str.split(","):
                    tgt = tgt.strip()
                    if tgt:
                        transitions.append([st, tgt])
            sm = StateMachine(states=states, initial=initial, labeling=labeling, transitions=transitions)

    safety_props: list[str] = []
    if "ThuocTinhAnToan" in wb.sheetnames:
        rows = list(wb["ThuocTinhAnToan"].iter_rows(values_only=True))
        if len(rows) > 1:
            for r in rows[1:]:
                if r and r[0]:
                    safety_props.append(str(r[0]).strip())

    mutual_ex: dict[str, str] | None = None
    if "LoaiTruLanNhau" in wb.sheetnames:
        rows = list(wb["LoaiTruLanNhau"].iter_rows(values_only=True))
        if len(rows) > 1 and rows[1] and rows[1][0] and rows[1][1]:
            mutual_ex = {"atom1": str(rows[1][0]).strip(), "atom2": str(rows[1][1]).strip()}

    funcs: list[FunctionSpec] = []
    if "HamSinhMa" in wb.sheetnames:
        rows = list(wb["HamSinhMa"].iter_rows(values_only=True))
        if len(rows) > 1:
            for r in rows[1:]:
                if not r or not r[0]:
                    continue
                fname = str(r[0]).strip()
                lang = str(r[1]).strip().lower() if len(r) > 1 and r[1] else "c"
                kind = str(r[2]).strip() if len(r) > 2 and r[2] else "buffer_copy"
                params_str = str(r[3]).strip() if len(r) > 3 and r[3] else ""
                params = [p.strip() for p in params_str.split(",") if p.strip()]
                buf_sz = int(r[4]) if len(r) > 4 and r[4] is not None and str(r[4]).isdigit() else None
                bug = str(r[5]).strip() if len(r) > 5 and r[5] else None
                if bug in ("none", "null", ""):
                    bug = None
                mv = []
                if len(r) > 6 and r[6] is not None:
                    try:
                        mv.append(int(r[6]))
                    except ValueError:
                        pass
                if len(r) > 7 and r[7] is not None:
                    try:
                        mv.append(int(r[7]))
                    except ValueError:
                        pass
                funcs.append(
                    FunctionSpec(
                        name=fname,
                        language=lang,
                        kind=kind,
                        params=params,
                        buffer_size=buf_sz,
                        bug=bug,
                        magic_values=mv,
                    )
                )

    checks: list[ArithmeticCheck] = []
    if "KiemTraSoHoc" in wb.sheetnames:
        rows = list(wb["KiemTraSoHoc"].iter_rows(values_only=True))
        if len(rows) > 1:
            for r in rows[1:]:
                if not r or not r[0]:
                    continue
                cname = str(r[0]).strip()
                vars_str = str(r[1]).strip() if len(r) > 1 and r[1] else ""
                ivars = [v.strip() for v in vars_str.split(",") if v.strip()]
                cond = str(r[2]).strip() if len(r) > 2 and r[2] else None
                ovf = str(r[3]).strip() if len(r) > 3 and r[3] else None
                checks.append(ArithmeticCheck(name=cname, int_vars=ivars, condition=cond, overflow_expr=ovf))

    tests: list[TestCase] = []
    if "TestCase" in wb.sheetnames:
        rows = list(wb["TestCase"].iter_rows(values_only=True))
        if len(rows) > 1:
            for r in rows[1:]:
                if not r or not r[0]:
                    continue
                tid = str(r[0]).strip()
                req_ids_str = str(r[1]).strip() if len(r) > 1 and r[1] else ""
                rids = [x.strip() for x in req_ids_str.split(",") if x.strip()]
                title = str(r[2]).strip() if len(r) > 2 and r[2] else ""
                pre = str(r[3]).strip() if len(r) > 3 and r[3] else ""
                steps_str = str(r[4]).strip() if len(r) > 4 and r[4] else ""
                steps = [s.strip() for s in steps_str.split("\n") if s.strip()]
                exp = str(r[5]).strip() if len(r) > 5 and r[5] else ""
                prio = str(r[6]).strip() if len(r) > 6 and r[6] else "HIGH"
                tests.append(
                    TestCase(
                        id=tid,
                        requirement_ids=rids,
                        title=title,
                        preconditions=pre,
                        steps=steps,
                        expected_result=exp,
                        priority=prio,
                    )
                )

    return Specification(
        name=name,
        description=description,
        requirements=reqs,
        state_machine=sm,
        additional_flows={},
        safety_properties=safety_props,
        mutual_exclusion=mutual_ex,
        functions=funcs,
        arithmetic_checks=checks,
        test_cases=tests,
        source_file=str(path),
    )


def load_specification(filepath: str | Path) -> Specification:
    path = Path(filepath)
    if path.suffix.lower() in (".xlsx", ".xls"):
        return parse_specification_excel(path)
    return parse_specification_json(path)
